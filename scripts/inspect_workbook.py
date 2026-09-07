from __future__ import annotations

import json
from pathlib import Path

import openpyxl

WORKBOOK = Path(__file__).resolve().parents[1] / "三年级英语词汇_Master_Database_v1.xlsx"


def main() -> None:
    workbook = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    report: dict[str, object] = {"path": str(WORKBOOK), "sheets": []}
    for worksheet in workbook.worksheets:
        rows = worksheet.iter_rows(values_only=True)
        header = list(next(rows, ()))
        samples = [list(row) for _, row in zip(range(3), rows)]
        report["sheets"].append(
            {
                "name": worksheet.title,
                "max_row": worksheet.max_row,
                "max_column": worksheet.max_column,
                "header": header,
                "samples": samples,
            }
        )
    print(json.dumps(report, ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
